"""Archivos descargables: Excel (openpyxl) y PDF (fpdf2)."""

import io
from datetime import datetime

from flask import send_file
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

COLOR_NUMI = "0F9D8A"


def enviar(contenido, nombre, tipo):
    return send_file(io.BytesIO(contenido), mimetype=tipo, as_attachment=True, download_name=nombre)


def excel(hojas):
    """hojas = [(titulo, encabezados, filas), ...] → bytes de un .xlsx"""
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, encabezados, filas in hojas:
        hoja = libro.create_sheet(titulo[:31])  # Excel no acepta títulos de más de 31 caracteres
        hoja.append(encabezados)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=COLOR_NUMI)
        for fila in filas:
            hoja.append(list(fila))
        for i, encabezado in enumerate(encabezados, start=1):
            largo = max([len(str(encabezado))] + [len(str(f[i - 1])) for f in filas])
            hoja.column_dimensions[hoja.cell(1, i).column_letter].width = min(45, largo + 3)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _latin(texto):
    # Las fuentes básicas del PDF solo tienen caracteres latinos (tildes y ñ sí)
    return str(texto).replace("—", "-").encode("latin-1", "replace").decode("latin-1")


def pdf(titulo, subtitulo, secciones):
    """secciones = [(titulo, encabezados, filas), ...] → bytes de un .pdf"""
    doc = FPDF(format="A4")
    doc.set_auto_page_break(auto=True, margin=15)
    doc.add_page()
    doc.set_font("Helvetica", "B", 16)
    doc.set_text_color(15, 157, 138)
    doc.cell(0, 10, _latin(titulo), new_x="LMARGIN", new_y="NEXT")
    doc.set_font("Helvetica", "", 10)
    doc.set_text_color(90, 90, 90)
    doc.multi_cell(0, 5, _latin(f"{subtitulo}\nGenerado el {datetime.now():%d/%m/%Y %H:%M} · Numi"))
    doc.set_text_color(0, 0, 0)

    for nombre, encabezados, filas in secciones:
        doc.ln(4)
        doc.set_font("Helvetica", "B", 12)
        doc.cell(0, 8, _latin(nombre), new_x="LMARGIN", new_y="NEXT")
        doc.set_font("Helvetica", "", 9)
        if not filas:
            doc.cell(0, 6, "Sin datos para los filtros elegidos.", new_x="LMARGIN", new_y="NEXT")
            continue
        with doc.table(text_align="LEFT", line_height=6) as tabla:
            fila = tabla.row()
            for e in encabezados:
                fila.cell(_latin(e))
            for datos in filas:
                fila = tabla.row()
                for valor in datos:
                    fila.cell(_latin("" if valor is None else valor))
    return bytes(doc.output())
